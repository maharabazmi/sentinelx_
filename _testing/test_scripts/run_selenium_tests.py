import os
import sys
import unittest
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as conditions
from selenium.webdriver.support.ui import WebDriverWait


TESTING_DIR = Path(__file__).resolve().parents[1]
BASE_URL = os.getenv("SENTINELX_BASE_URL", "http://127.0.0.1:3000").rstrip("/")
HEADLESS = os.getenv("SENTINELX_HEADLESS", "true").lower() not in {"0", "false", "no"}
BROWSER = os.getenv("SENTINELX_BROWSER", "chrome").lower()

TEST_CASES = [
    {
        "id": "UI-001",
        "method": "test_landing_page_loads",
        "module": "Public landing",
        "title": "Landing page renders its public safety content",
        "preconditions": "Frontend development server is running.",
        "steps": "Open the base URL and wait for the landing page heading.",
        "expected": "The public defense heading and official sign-in action are visible.",
    },
    {
        "id": "UI-002",
        "method": "test_emergency_hotlines_dropdown",
        "module": "Emergency hotline",
        "title": "Hotline menu shows Bangladesh emergency contacts",
        "preconditions": "Landing page is loaded.",
        "steps": "Open the National Emergency Hotlines menu.",
        "expected": "The menu shows 999, 16121, 109, and 333.",
    },
    {
        "id": "UI-003",
        "method": "test_login_modal_and_password_visibility",
        "module": "Authentication UI",
        "title": "Login modal opens and password visibility toggles",
        "preconditions": "Landing page is loaded; no account credentials are needed.",
        "steps": "Open Official Sign In, reveal the password, then hide it again.",
        "expected": "The login dialog opens and the password input toggles text/password types.",
    },
    {
        "id": "UI-004",
        "method": "test_registration_modal_opens",
        "module": "Registration UI",
        "title": "NID registration wizard opens from the landing page",
        "preconditions": "Landing page is loaded; no registration is submitted.",
        "steps": "Select Verify NID & Enter.",
        "expected": "The registration wizard's identity-verification step is visible.",
    },
    {
        "id": "UI-005",
        "method": "test_theme_toggle",
        "module": "Display settings",
        "title": "Theme control switches between light and dark",
        "preconditions": "Landing page is loaded.",
        "steps": "Read the theme control label, click it, and read the updated label.",
        "expected": "The control changes from Switch to light theme to Switch to dark theme or vice versa.",
    },
]


def create_driver():
    if BROWSER == "edge":
        options = webdriver.EdgeOptions()
        if HEADLESS:
            options.add_argument("--headless=new")
        options.add_argument("--window-size=1440,1000")
        return webdriver.Edge(options=options)

    if BROWSER != "chrome":
        raise ValueError("SENTINELX_BROWSER must be 'chrome' or 'edge'.")

    options = webdriver.ChromeOptions()
    if HEADLESS:
        options.add_argument("--headless=new")
    options.add_argument("--window-size=1440,1000")
    options.add_argument("--disable-dev-shm-usage")
    return webdriver.Chrome(options=options)


class SentinelXSmokeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.driver = create_driver()
        cls.wait = WebDriverWait(cls.driver, 15)

    @classmethod
    def tearDownClass(cls):
        if hasattr(cls, "driver"):
            cls.driver.quit()

    def setUp(self):
        self.driver.get(BASE_URL)
        self.wait.until(conditions.presence_of_element_located((By.TAG_NAME, "body")))

    def test_landing_page_loads(self):
        self.wait.until(conditions.visibility_of_element_located((By.XPATH, "//h1[contains(., 'PUBLIC DEFENSE')]")))
        self.assertTrue(self.driver.find_element(By.XPATH, "//button[.//span[contains(., 'OFFICIAL SIGN IN')]]").is_displayed())

    def test_emergency_hotlines_dropdown(self):
        self.driver.find_element(By.CSS_SELECTOR, "button[title='National Emergency Hotlines']").click()
        menu = self.wait.until(conditions.visibility_of_element_located((
            By.XPATH,
            "//h4[contains(., 'Bangladesh National Hotlines')]/ancestor::div[contains(@class, 'w-72')][1]",
        )))
        menu_text = menu.text
        for hotline in ("999", "16121", "109", "333"):
            self.assertIn(hotline, menu_text)

    def test_login_modal_and_password_visibility(self):
        self.driver.find_element(By.XPATH, "//button[.//span[contains(., 'OFFICIAL SIGN IN')]]").click()
        self.wait.until(conditions.visibility_of_element_located((
            By.XPATH,
            "//*[contains(., 'Sign In to SentinelX')]",
        )))
        password = self.driver.find_element(By.CSS_SELECTOR, "form input[type='password']")
        self.driver.find_element(By.CSS_SELECTOR, "button[aria-label='Show password']").click()
        self.wait.until(lambda driver: password.get_attribute("type") == "text")
        self.driver.find_element(By.CSS_SELECTOR, "button[aria-label='Hide password']").click()
        self.wait.until(lambda driver: password.get_attribute("type") == "password")
        self.assertEqual(password.get_attribute("type"), "password")

    def test_registration_modal_opens(self):
        self.driver.find_element(By.XPATH, "//button[.//span[contains(., 'VERIFY NID & ENTER')]]").click()
        self.wait.until(conditions.visibility_of_element_located((
            By.XPATH,
            "//*[contains(., 'Citizen Identity Verification')]",
        )))

    def test_theme_toggle(self):
        theme_button = self.driver.find_element(By.CSS_SELECTOR, "button[aria-label^='Switch to']")
        initial_label = theme_button.get_attribute("aria-label")
        theme_button.click()
        self.wait.until(lambda driver: driver.find_element(
            By.CSS_SELECTOR,
            "button[aria-label^='Switch to']",
        ).get_attribute("aria-label") != initial_label)
        updated_label = self.driver.find_element(By.CSS_SELECTOR, "button[aria-label^='Switch to']").get_attribute("aria-label")
        self.assertNotEqual(initial_label, updated_label)


class ResultCollector(unittest.TextTestResult):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.outcomes = {}

    def addSuccess(self, test):
        super().addSuccess(test)
        self.outcomes[test._testMethodName] = ("PASS", "")

    def addFailure(self, test, err):
        super().addFailure(test, err)
        self.outcomes[test._testMethodName] = ("FAIL", self._exc_info_to_string(err, test))

    def addError(self, test, err):
        super().addError(test, err)
        self.outcomes[test._testMethodName] = ("ERROR", self._exc_info_to_string(err, test))

    def addSkip(self, test, reason):
        super().addSkip(test, reason)
        self.outcomes[test._testMethodName] = ("SKIP", reason)


def write_artifacts(outcomes, started_at, total, failures, errors):
    rows = []
    for case in TEST_CASES:
        status, details = outcomes.get(case["method"], ("NOT RUN", ""))
        rows.append({**case, "status": status, "actual": details or status, "defect": details if status in {"FAIL", "ERROR"} else ""})

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Test Cases"
    headers = ["Test Case ID", "Module", "Test Title", "Preconditions", "Test Steps", "Expected Result", "Actual Result", "Status", "Defect / Feedback"]
    sheet.append(headers)
    for row in rows:
        sheet.append([row[key] for key in ("id", "module", "title", "preconditions", "steps", "expected", "actual", "status", "defect")])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
    for column, width in enumerate((15, 22, 42, 42, 48, 48, 48, 14, 48), start=1):
        sheet.column_dimensions[chr(64 + column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(TESTING_DIR / "test_cases.xlsx")

    report_path = TESTING_DIR / "test_report.pdf"
    document = SimpleDocTemplate(str(report_path), pagesize=landscape(A4), rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("SentinelX Selenium UI Test Report", styles["Title"]),
        Paragraph(f"Run started: {started_at.isoformat(timespec='seconds')} UTC | Target: {BASE_URL} | Browser: {BROWSER}", styles["BodyText"]),
        Paragraph(f"Summary: {total} executed | {len([row for row in rows if row['status'] == 'PASS'])} passed | {failures} failed | {errors} errors", styles["BodyText"]),
        Spacer(1, 12),
    ]
    table_data = [["ID", "Test", "Expected", "Status", "Actual / Defect"]]
    for row in rows:
        table_data.append([
            row["id"],
            Paragraph(row["title"], styles["BodyText"]),
            Paragraph(row["expected"], styles["BodyText"]),
            row["status"],
            Paragraph(row["actual"], styles["BodyText"]),
        ])
    table = Table(table_data, colWidths=(42, 145, 185, 55, 315), repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#9EADBC")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(table)
    document.build(story)


def main():
    started_at = datetime.now(timezone.utc).replace(tzinfo=None)
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(SentinelXSmokeTests)
    runner = unittest.TextTestRunner(resultclass=ResultCollector, verbosity=2)
    result = runner.run(suite)
    write_artifacts(result.outcomes, started_at, result.testsRun, len(result.failures), len(result.errors))
    print(f"Artifacts written to {TESTING_DIR}")
    return 0 if result.wasSuccessful() else 1


if __name__ == "__main__":
    sys.exit(main())